"""Copy one range out of Mark's workbook into a tiny spreadsheet of its own.

**This is the piece that keeps Excel away from his files.** Nothing here starts
Office. It opens the workbook with openpyxl, reads it, and writes a small new
`.xlsx` somewhere disposable. Whatever renders the picture afterwards opens only
that copy, so no mechanism in this app can reach one of his workbooks with a
program that is able to save.

That matters more than it sounds. `HOW-WE-WORK.md` says never write into
`Report Examples/`, and until now that was a rule a person had to keep. Here it
becomes a thing the design cannot do.

Three practical reasons on top of the safety one, all measured in the locker
before this was written:

- His workbooks are macro files and several are over 10 MB. Opening one in
  Excel is slow and prompts.
- A merged banner styled anywhere but its top-left cell trips Excel's repair
  dialog. Copying only the top-left cell's style avoids it.
- Rows and columns hidden in the source are dropped, which is what his own
  copy-and-paste into Word already does.

**Colours are resolved to literal red-green-blue here, on purpose.** A cell can
name its colour by an index into the workbook's own palette, or by a theme slot
plus a lightening amount. The new little workbook has neither a custom palette
nor the theme, so an index that means near-white in his workbook means tan in
the standard one. That was a real fault, found by looking at an export in July
2026: whole bands of a grid came out orange. Every colour is turned into a
literal value while the source is still open to be asked.

Adapted from the locker's `xlsm_exhibit.py`, which is quarry and not a drop-in.

openpyxl only. No Office, no network, and nothing platform-specific, so all of
this runs and is tested on any machine.
"""
import copy as copylib
import datetime
import re
import zipfile
from pathlib import Path

# Kept narrow on purpose. A column much narrower than this renders as a sliver
# and a column much wider pushes the grid off the page, and neither looks like
# the delivered reports. Measured in the locker.
MIN_COLUMN_WIDTH = 11.5
MAX_COLUMN_WIDTH = 16.0

DEFAULT_COLUMN_WIDTH = 8.43


class CannotRead(Exception):
    """The workbook or the range is not what was expected.

    Worded the way Mark or Colleen should read it, like `RevealFailed` and
    `InstallRefused` elsewhere in this app. Nothing here ever reaches a person
    as a traceback.
    """

    def __init__(self, message: str):
        super().__init__(message)
        self.message = message


# ------------------------------------------------------------- colours -----
def _indexed_palette(workbook_path: Path) -> list:
    """The workbook's own legacy colour list, or the standard one.

    Read straight out of the file's `styles.xml` rather than through openpyxl,
    because openpyxl gives back the standard list when a workbook carries its
    own. Index 47 is near-white in Mark's palette and tan in the standard one,
    which is the whole reason this function exists.
    """
    try:
        with zipfile.ZipFile(str(workbook_path)) as bundle:
            styles = bundle.read("xl/styles.xml").decode("utf-8", "replace")
        found = re.search(r"<indexedColors>(.*?)</indexedColors>", styles, re.S)
        if found:
            listed = re.findall(r'rgb="([0-9A-Fa-f]{8})"', found.group(1))
            if listed:
                return listed
    except (OSError, KeyError, zipfile.BadZipFile, ValueError):
        # An unreadable style sheet is not a reason to fail: the standard
        # palette below is right for most workbooks and wrong only in colour.
        pass
    from openpyxl.styles.colors import COLOR_INDEX
    return list(COLOR_INDEX)


def _theme_palette(source_workbook) -> list:
    """The ten theme colours, in the order a cell's theme number counts them."""
    xml = getattr(source_workbook, "loaded_theme", None)
    if not xml:
        return []
    if isinstance(xml, bytes):
        xml = xml.decode("utf-8", "replace")
    scheme = re.search(r"<a:clrScheme.*?</a:clrScheme>", xml, re.S)
    if not scheme:
        return []
    slots = {}
    pattern = (r'<a:(\w+)>\s*<a:(?:sysClr val="\w+" lastClr="([0-9A-Fa-f]{6})"'
               r'|srgbClr val="([0-9A-Fa-f]{6})")')
    for name, system_colour, plain_colour in re.findall(pattern, scheme.group(0)):
        slots[name] = system_colour or plain_colour
    order = ("lt1", "dk1", "lt2", "dk2", "accent1", "accent2", "accent3",
             "accent4", "accent5", "accent6")
    return [slots.get(name) for name in order]


def _tinted(colour: str, tint: float) -> str:
    """Lighten or darken one colour by the amount a theme cell asks for."""
    def one(value):
        if tint > 0:
            value = value + (255 - value) * tint
        elif tint < 0:
            value = value * (1 + tint)
        return max(0, min(255, int(round(value))))

    red, green, blue = (int(colour[at:at + 2], 16) for at in (0, 2, 4))
    return "%02X%02X%02X" % (one(red), one(green), one(blue))


def _literal_colour(colour, palette: list, theme: list):
    """A plain colour value for one that is named indirectly, else None.

    None means keep what is already there, which covers a colour that is
    already literal and a colour we cannot work out. Guessing would be worse
    than leaving it: a wrong literal is permanent, and an unresolved one at
    least renders the way the standard palette says.
    """
    try:
        if colour is None or colour.type == "rgb":
            return None
        if colour.type == "indexed" and palette:
            index = colour.indexed
            if isinstance(index, int) and 0 <= index < len(palette):
                return "FF" + palette[index][-6:]
        if colour.type == "theme" and theme:
            slot = colour.theme
            if isinstance(slot, int) and 0 <= slot < len(theme) and theme[slot]:
                return "FF" + _tinted(theme[slot], colour.tint or 0.0)
    except (AttributeError, TypeError, ValueError, IndexError):
        return None
    return None


# -------------------------------------------------------------- reading ----
def _open_source(workbook: Path):
    """Open his workbook to be read and never to be written.

    `data_only=True` asks for what a formula last worked out rather than the
    formula, which is what a picture of a grid has to show. openpyxl never
    writes to a file it loads, and nothing in this module calls save on it.
    """
    workbook = Path(workbook)
    if not workbook.is_file():
        raise CannotRead("There is no workbook at:\n    %s" % workbook)
    try:
        import openpyxl
    except ImportError:
        raise CannotRead(
            "This copy of the app cannot read spreadsheets, which means it was "
            "built without openpyxl. Send Spenser this message.")
    try:
        return openpyxl.load_workbook(str(workbook), data_only=True)
    except Exception as exc:
        raise CannotRead(
            "That workbook could not be read:\n    %s\n\n%s\n"
            "Nothing was changed." % (workbook, exc))


def sheet_names(workbook: Path) -> list:
    """What sheets a workbook holds. For telling somebody what to ask for."""
    return list(_open_source(workbook).sheetnames)


def extract(workbook: Path, sheet: str, cell_range: str, into: Path,
            landscape: bool = False) -> Path:
    """Write the range into a small standalone `.xlsx` and return its path.

    `into` is a folder the caller owns and cleans up. Deliberately not a
    temporary folder made here: this is called in a chain of steps and the
    caller keeps every intermediate file when a person is trying to work out
    what went wrong.
    """
    import openpyxl
    from openpyxl.styles.colors import Color
    from openpyxl.utils import get_column_letter, range_boundaries

    source = _open_source(workbook)
    if sheet not in source.sheetnames:
        raise CannotRead(
            "That workbook has no sheet called %r.\n\n"
            "It has these:\n    %s" % (sheet, "\n    ".join(source.sheetnames)))
    from_sheet = source[sheet]

    try:
        first_col, first_row, last_col, last_row = range_boundaries(cell_range)
    except (ValueError, TypeError):
        raise CannotRead(
            "%r is not a range of cells. A range looks like AY98:BK102."
            % cell_range)
    if None in (first_col, first_row, last_col, last_row):
        raise CannotRead(
            "%r is not a whole range of cells. A range looks like AY98:BK102."
            % cell_range)

    def column_is_hidden(index):
        found = from_sheet.column_dimensions.get(get_column_letter(index))
        return bool(found and found.hidden)

    def row_is_hidden(index):
        found = from_sheet.row_dimensions.get(index)
        return bool(found and found.hidden)

    # Hidden rows and columns are dropped rather than carried across, because
    # that is what his own paste into Word shows. The assessment block hides
    # its empty parcel slots so they collapse.
    source_cols = [c for c in range(first_col, last_col + 1) if not column_is_hidden(c)]
    source_rows = [r for r in range(first_row, last_row + 1) if not row_is_hidden(r)]
    if not source_cols or not source_rows:
        raise CannotRead(
            "Every row or column in %s on sheet %r is hidden, so there is "
            "nothing to show." % (cell_range, sheet))

    col_at = dict((c, i + 1) for i, c in enumerate(source_cols))
    row_at = dict((r, i + 1) for i, r in enumerate(source_rows))

    little = openpyxl.Workbook()
    # The theme travels too, as a second line of defence behind resolving every
    # colour to a literal below.
    loaded_theme = getattr(source, "loaded_theme", None)
    if loaded_theme:
        try:
            little.loaded_theme = loaded_theme
        except (AttributeError, TypeError):
            pass
    to_sheet = little.active
    to_sheet.title = sheet[:31]

    palette = _indexed_palette(workbook)
    theme = _theme_palette(source)

    default_width = from_sheet.sheet_format.defaultColWidth or DEFAULT_COLUMN_WIDTH
    for column in source_cols:
        found = from_sheet.column_dimensions.get(get_column_letter(column))
        width = found.width if (found and found.width) else default_width
        width = min(max(width, MIN_COLUMN_WIDTH), MAX_COLUMN_WIDTH)
        to_sheet.column_dimensions[get_column_letter(col_at[column])].width = width

    merged = list(from_sheet.merged_cells.ranges)

    def is_covered_but_not_the_corner(column, row):
        """Part of a merged block, and not the cell that carries its style."""
        for block in merged:
            if (block.min_row <= row <= block.max_row
                    and block.min_col <= column <= block.max_col):
                return not (row == block.min_row and column == block.min_col)
        return False

    for row in source_rows:
        target_row = row_at[row]
        found = from_sheet.row_dimensions.get(row)
        if found and found.height:
            to_sheet.row_dimensions[target_row].height = found.height
        for column in source_cols:
            here = from_sheet.cell(row=row, column=column)
            there = to_sheet.cell(row=target_row, column=col_at[column])

            value = here.value
            if isinstance(value, (datetime.datetime, datetime.date)):
                # Baked as text in the shape the delivered reports use. What a
                # foreign renderer does with a real date is not reliable, and a
                # date is one of the few things a reader spots instantly.
                there.value = "%d/%d/%d" % (value.month, value.day, value.year)
            else:
                there.value = value

            if is_covered_but_not_the_corner(column, row):
                # Styling a covered cell is what trips Excel's repair dialog.
                continue

            font = copylib.copy(here.font)
            literal = _literal_colour(font.color, palette, theme)
            if literal:
                font.color = Color(rgb=literal)
            there.font = font

            fill = copylib.copy(here.fill)
            try:
                if fill.patternType:
                    front = _literal_colour(fill.fgColor, palette, theme)
                    if front:
                        fill.fgColor = Color(rgb=front)
                    back = _literal_colour(fill.bgColor, palette, theme)
                    if back:
                        fill.bgColor = Color(rgb=back)
            except (AttributeError, TypeError, ValueError):
                pass
            there.fill = fill

            border = copylib.copy(here.border)
            try:
                # Border colours name themselves the same indirect way fills
                # do. Missing this turned Mark's grey row rules light green.
                for edge in ("left", "right", "top", "bottom", "diagonal"):
                    side = getattr(border, edge, None)
                    if side is not None and side.style:
                        literal = _literal_colour(side.color, palette, theme)
                        if literal:
                            side.color = Color(rgb=literal)
            except (AttributeError, TypeError, ValueError):
                pass
            there.border = border

            there.alignment = copylib.copy(here.alignment)
            there.number_format = here.number_format

    for block in merged:
        if block.max_row < first_row or block.min_row > last_row:
            continue
        if block.max_col < first_col or block.min_col > last_col:
            continue
        columns = [col_at[c] for c
                   in range(max(block.min_col, first_col), min(block.max_col, last_col) + 1)
                   if c in col_at]
        rows = [row_at[r] for r
                in range(max(block.min_row, first_row), min(block.max_row, last_row) + 1)
                if r in row_at]
        if not columns or not rows:
            continue
        if len(columns) < 2 and len(rows) < 2:
            # One visible cell left of a merge is not a merge any more, and
            # openpyxl refuses a single-cell one.
            continue
        to_sheet.merge_cells(start_row=min(rows), start_column=min(columns),
                             end_row=max(rows), end_column=max(columns))

    to_sheet.print_area = "A1:%s%d" % (get_column_letter(len(source_cols)),
                                       len(source_rows))
    to_sheet.page_setup.orientation = "landscape" if landscape else "portrait"
    to_sheet.page_setup.fitToWidth = 1
    to_sheet.page_setup.fitToHeight = 1
    from openpyxl.worksheet.properties import PageSetupProperties
    to_sheet.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)

    into = Path(into)
    into.mkdir(parents=True, exist_ok=True)
    written = into / "grid.xlsx"
    little.save(str(written))
    return written
