"""Run the whole Office chain once and say what happened at each step.

**This is how anybody finds out whether the app can drive Excel and Word on a
particular computer.** It is the proof for Phase 1, and afterwards it is a
support tool: when something Office-shaped goes wrong in Mark's office, one
line run there answers the question without anybody guessing.

It ships inside the package on purpose, decided by Spenser on 2026-09-04. No
screen and no button reaches it, which is the same arrangement
`install_windows.py` and `update_apply.py` already have.

Two ways to run it.

    office_check.py

With nothing after it, it builds its own small spreadsheet and works on that.
That answers the narrow question, which is whether Excel and Word can be driven
here at all, and it needs no material and no setup.

    office_check.py <workbook> <sheet> <range>

Given one of Mark's workbooks it does the same thing to a real grid. That is the
acceptance run, and it is the one `docs/ROADMAP.md` asks for.

**Every file it makes goes in one folder in the person's own home folder**, and
nothing is written near a job. Working files are kept rather than tidied away,
because the moment anybody runs this is the moment they want to see where the
chain stopped.

It never raises. A failure is a sentence and an exit code, because the person
reading it is not going to debug a traceback.
"""
import os
import shutil
import sys
import traceback
from pathlib import Path

HERE = Path(__file__).resolve().parent
if str(HERE) not in sys.path:
    sys.path.insert(0, str(HERE))

import gridextract  # noqa: E402
import office  # noqa: E402

FOLDER_NAME = "Roy R. Fisher office check"


def where() -> Path:
    """One folder, in the person's own home, easy to say out loud.

    RRF_OFFICE_CHECK overrides it, the same way RRF_INSTALL_HOME and
    RRF_KEY_FILE already do elsewhere. Tests need it: this folder is cleared at
    the start of every run, and a test that used the real one would delete
    whatever a person had left there.
    """
    override = os.environ.get("RRF_OFFICE_CHECK")
    return Path(override) if override else Path.home() / FOLDER_NAME


def a_practice_workbook(into: Path) -> Path:
    """A small grid shaped like one of Mark's, built here rather than shipped.

    Built rather than carried as a file so there is nothing to keep in step
    with anything. It has a merged heading, a filled header row, borders and
    numbers, which are the four things that go wrong when a grid renders badly.

    **The headings wrap, because Mark's do.** Measured 2026-09-04 against the
    Utica Ridge assessment grid. Without wrap a long heading overflows its cell
    and is cut at the edge of the copied range, which reads as a fault in the
    app and is a fault in this workbook. A practice grid that does not look
    like his material tests the wrong thing.
    """
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Practice"

    sheet.cell(row=1, column=1).value = "Real Estate Assessment and Taxes"
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=4)
    sheet.cell(row=1, column=1).font = Font(bold=True, size=12)
    sheet.cell(row=1, column=1).alignment = Alignment(horizontal="center")

    headings = ("Tax ID", "Year", "Land", "Total Assessment")
    edge = Side(style="thin", color="FF444444")
    box = Border(left=edge, right=edge, top=edge, bottom=edge)
    for column, heading in enumerate(headings, start=1):
        cell = sheet.cell(row=2, column=column)
        cell.value = heading
        cell.font = Font(bold=True)
        cell.fill = PatternFill(patternType="solid", fgColor="FFD9D9D9")
        cell.border = box
        cell.alignment = Alignment(horizontal="center", wrap_text=True)

    rows = (("Y0917-12J", 2025, 1021630, 4773560),
            ("Y0917-12J", 2024, 1021630, 4528110),
            ("Y0917-12J", 2023, 1021630, 4528110))
    for offset, values in enumerate(rows):
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row=3 + offset, column=column)
            cell.value = value
            cell.border = box
            if column > 2:
                cell.number_format = "#,##0"

    into.mkdir(parents=True, exist_ok=True)
    made = into / "practice.xlsx"
    book.save(str(made))
    return made


def a_document_holding(picture: Path, out_docx: Path) -> Path:
    """One Word file with the grid in it, the way a section will carry one."""
    from docx import Document
    from docx.shared import Inches

    document = Document()
    document.add_paragraph("Roy R. Fisher: Office check")
    document.add_paragraph(
        "This document was made by the app to prove it can drive Excel and "
        "Word on this computer. It is not part of any report.")
    document.add_picture(str(picture), width=Inches(6.0))
    document.save(str(out_docx))
    return out_docx


def _size(path: Path) -> str:
    try:
        return "%.1f KB" % (path.stat().st_size / 1024.0)
    except OSError:
        return "not there"


def run(workbook=None, sheet=None, cell_range=None, out=print) -> int:
    """The whole chain. Returns 0 when every step worked."""
    place = where()
    shutil.rmtree(str(place), ignore_errors=True)
    place.mkdir(parents=True, exist_ok=True)

    out("Roy R. Fisher: can this computer drive Excel and Word?")
    out("")
    out("Everything it makes goes here:")
    out("    %s" % place)
    out("")

    out("1. What is going to do the work")
    out("   %s" % office.describe())
    out("")

    if workbook is None:
        out("2. Making a small spreadsheet to practise on")
        try:
            workbook = a_practice_workbook(place)
        except Exception as exc:
            out("   It could not be made: %s" % exc)
            return 1
        sheet, cell_range = "Practice", "A1:D5"
        out("   %s" % workbook.name)
    else:
        workbook = Path(workbook)
        out("2. Using the workbook you named")
        out("   %s" % workbook)
        if sheet is None or cell_range is None:
            out("")
            out("   A sheet and a range are needed too. For example:")
            out("       office_check.py \"%s\" Assessment AY98:BK102" % workbook.name)
            try:
                out("")
                out("   That workbook holds these sheets:")
                for one in gridextract.sheet_names(workbook):
                    out("       %s" % one)
            except gridextract.CannotRead as exc:
                out("   %s" % exc.message)
            return 1
    out("")

    out("3. Taking %s off sheet %r, without opening Excel" % (cell_range, sheet))
    picture = place / "grid.png"
    try:
        office.grid_to_image(workbook, sheet, cell_range, picture,
                             workshop=place)
    except (office.OfficeRefused, gridextract.CannotRead) as exc:
        out("")
        out(exc.message)
        return 1
    except Exception:
        out("")
        out("   Something nobody expected went wrong. All of this is useful:")
        out(traceback.format_exc())
        return 1
    out("   grid.xlsx   %s   the small copy Excel opened" % _size(place / "grid.xlsx"))
    out("   grid.png    %s   the picture of the grid" % _size(picture))
    out("")

    out("4. Putting that picture in a Word document")
    document = place / "office-check.docx"
    try:
        a_document_holding(picture, document)
    except Exception as exc:
        out("   The document could not be made: %s" % exc)
        return 1
    out("   office-check.docx   %s" % _size(document))
    out("")

    out("5. Asking Word for a PDF of it")
    pdf = place / "office-check.pdf"
    try:
        office.docx_to_pdf(document, pdf)
    except office.OfficeRefused as exc:
        out("")
        out(exc.message)
        return 1
    except Exception:
        out("")
        out("   Something nobody expected went wrong. All of this is useful:")
        out(traceback.format_exc())
        return 1
    out("   office-check.pdf    %s" % _size(pdf))
    out("")

    out("Everything worked. This computer can drive Excel and Word.")
    out("")
    out("Open the folder above and look at the PDF. It should hold the grid,")
    out("and the grid should look the way it does in Excel.")
    return 0


def main(argv=None) -> int:
    argv = list(sys.argv[1:] if argv is None else argv)
    if argv and argv[0] in ("-h", "--help"):
        print(__doc__)
        return 0
    workbook = argv[0] if len(argv) > 0 else None
    sheet = argv[1] if len(argv) > 1 else None
    cell_range = argv[2] if len(argv) > 2 else None
    return run(workbook, sheet, cell_range)


if __name__ == "__main__":
    sys.exit(main())
