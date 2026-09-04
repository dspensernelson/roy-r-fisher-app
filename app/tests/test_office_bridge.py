"""Taking a grid out of a workbook, and asking a computer to draw it.

Two halves, and the split is the whole design. Getting the range out is pure
Python and runs anywhere, so nearly all of it is tested here. Drawing the
picture needs Microsoft Office and is tested by hand, on a real machine, in
`docs/CHECKS.md`.

**Nothing here starts Excel or Word.** Every test that reaches the rendering
step hands in a stand-in, so this file behaves the same on a machine with
Office and a machine without it. A test that needed real Office would skip on
most machines, and `conftest.py` records twice what a quiet skip costs: it
reads as "this computer does not have it" and makes untested code look proven.

Most of the workbooks here are built in the test, so a clone with none of
Mark's material still proves the behaviour. One test reads a delivered workbook
and its only job is to prove the reading changes nothing.
"""
import sys
from pathlib import Path

import pytest

APP = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(APP / "engine"))
sys.path.insert(0, str(APP / "server"))

import gridextract  # noqa: E402
import office  # noqa: E402
import reveal  # noqa: E402

from conftest import CORPUS  # noqa: E402

openpyxl = pytest.importorskip("openpyxl", reason="openpyxl is a shipped dependency")

# One of Mark's delivered workbooks, and a range that holds his assessment
# grid. Named exactly, the way conftest names the template it needs, so a
# machine without the corpus says which piece is missing rather than skipping
# a whole file for a reason nobody reads.
REAL_WORKBOOK = (CORPUS / "DAVENPORT_5515 Utica Ridge - 2025 Tax"
                 / "DAVENPORT_5515 Utica Ridge Road-2025 Tax.xlsm")
REAL_SHEET = "Assessment"
REAL_RANGE = "AY98:BK102"

has_a_real_workbook = pytest.mark.skipif(
    not REAL_WORKBOOK.is_file(),
    reason="Mark's delivered workbooks are private and not in this repository",
)


# --------------------------------------------------------------- helpers ---
def a_workbook(path, rows=3, columns=3):
    """A small workbook with known contents, so a test can state its shape."""
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Grid"
    for row in range(1, rows + 1):
        for column in range(1, columns + 1):
            sheet.cell(row=row, column=column).value = "r%dc%d" % (row, column)
    book.save(str(path))
    return path


def values_in(path):
    """Every cell of the extracted workbook, as rows of values."""
    sheet = openpyxl.load_workbook(str(path)).active
    return [[cell.value for cell in row] for row in sheet.iter_rows()]


class FakeOffice:
    """Stands in for Excel and Word. Records what it was asked to do."""

    def __init__(self, here=True, writes=b"a picture"):
        self.here = here
        self.writes = writes
        self.grids = []
        self.documents = []

    def available(self):
        return self.here

    def name(self):
        return "a stand-in, not real Office"

    def render_grid(self, little, out_png):
        self.grids.append((Path(little), Path(out_png)))
        if self.writes is not None:
            Path(out_png).write_bytes(self.writes)

    def docx_to_pdf(self, docx, out_pdf):
        self.documents.append((Path(docx), Path(out_pdf)))
        if self.writes is not None:
            Path(out_pdf).write_bytes(self.writes)


# ------------------------------------------------- his files are not ours ---
@has_a_real_workbook
def test_reading_one_of_marks_workbooks_does_not_change_it(tmp_path):
    """The rule the whole design exists to make unbreakable.

    Byte for byte, not just the modified time. A workbook rewritten with the
    same contents would keep neither, but this is the check that would catch a
    save nobody meant to make.
    """
    before = REAL_WORKBOOK.read_bytes()
    stamp = REAL_WORKBOOK.stat().st_mtime

    gridextract.extract(REAL_WORKBOOK, REAL_SHEET, REAL_RANGE, tmp_path / "work")

    assert REAL_WORKBOOK.read_bytes() == before, "his workbook was written to"
    assert REAL_WORKBOOK.stat().st_mtime == stamp


@has_a_real_workbook
def test_a_real_grid_comes_out_with_its_numbers(tmp_path):
    written = gridextract.extract(REAL_WORKBOOK, REAL_SHEET, REAL_RANGE,
                                  tmp_path / "work")
    assert written.is_file()
    flat = [cell for row in values_in(written) for cell in row]
    assert "Real Estate Assessment and Taxes" in flat, "the grid's own title"
    assert "Total Assessment" in flat, "one of its column headings"


def test_the_extraction_never_opens_office():
    """It is the half that has to run everywhere, so it may not reach for
    Excel even indirectly."""
    source = Path(gridextract.__file__).read_text()
    for forbidden in ("subprocess", "osascript", "win32com", "pythoncom"):
        assert forbidden not in source, "gridextract reached for %s" % forbidden


def test_the_workbook_is_opened_for_reading_only():
    source = Path(gridextract.__file__).read_text()
    assert "data_only=True" in source
    assert "source.save" not in source


# ---------------------------------------------------- what comes across ----
def test_the_range_comes_across_at_the_right_size(tmp_path):
    book = a_workbook(tmp_path / "book.xlsx", rows=5, columns=5)
    written = gridextract.extract(book, "Grid", "B2:D4", tmp_path / "work")
    assert values_in(written) == [["r2c2", "r2c3", "r2c4"],
                                  ["r3c2", "r3c3", "r3c4"],
                                  ["r4c2", "r4c3", "r4c4"]]


def test_a_hidden_column_is_dropped_the_way_his_paste_drops_it(tmp_path):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Grid"
    for column, letter in enumerate("ABC", start=1):
        sheet.cell(row=1, column=column).value = letter
    sheet.column_dimensions["B"].hidden = True
    book.save(str(tmp_path / "book.xlsx"))

    written = gridextract.extract(tmp_path / "book.xlsx", "Grid", "A1:C1",
                                  tmp_path / "work")
    assert values_in(written) == [["A", "C"]], "the hidden column came across"


def test_a_hidden_row_is_dropped_too(tmp_path):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Grid"
    for row in (1, 2, 3):
        sheet.cell(row=row, column=1).value = "row%d" % row
    sheet.row_dimensions[2].hidden = True
    book.save(str(tmp_path / "book.xlsx"))

    written = gridextract.extract(tmp_path / "book.xlsx", "Grid", "A1:A3",
                                  tmp_path / "work")
    assert values_in(written) == [["row1"], ["row3"]]


def test_a_date_is_baked_as_the_text_the_reports_use(tmp_path):
    """A foreign renderer's idea of a date format is not reliable, and a date
    is one of the few things a reader spots instantly."""
    import datetime
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Grid"
    sheet.cell(row=1, column=1).value = datetime.date(2026, 3, 7)
    book.save(str(tmp_path / "book.xlsx"))

    written = gridextract.extract(tmp_path / "book.xlsx", "Grid", "A1:A1",
                                  tmp_path / "work")
    assert values_in(written) == [["3/7/2026"]]


def test_a_merged_heading_comes_across_merged(tmp_path):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Grid"
    sheet.cell(row=1, column=1).value = "One heading over three"
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=3)
    book.save(str(tmp_path / "book.xlsx"))

    written = gridextract.extract(tmp_path / "book.xlsx", "Grid", "A1:C1",
                                  tmp_path / "work")
    merged = openpyxl.load_workbook(str(written)).active.merged_cells.ranges
    assert [str(one) for one in merged] == ["A1:C1"]


def test_a_colour_named_by_number_becomes_a_real_colour(tmp_path):
    """The orange-bands fault. A cell can name its fill by a number that means
    one colour in Mark's workbook and a different one in a fresh workbook, so
    every colour is turned into a literal while his file is still open."""
    from openpyxl.styles import PatternFill
    from openpyxl.styles.colors import Color

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Grid"
    cell = sheet.cell(row=1, column=1)
    cell.value = "banded"
    cell.fill = PatternFill(patternType="solid", fgColor=Color(indexed=47))
    book.save(str(tmp_path / "book.xlsx"))

    written = gridextract.extract(tmp_path / "book.xlsx", "Grid", "A1:A1",
                                  tmp_path / "work")
    fill = openpyxl.load_workbook(str(written)).active.cell(row=1, column=1).fill
    assert fill.fgColor.type == "rgb", "the colour is still named by number"
    assert fill.fgColor.rgb not in (None, "00000000")


# ------------------------------------------------- saying what went wrong ---
def test_a_missing_workbook_says_so_in_a_sentence(tmp_path):
    with pytest.raises(gridextract.CannotRead) as caught:
        gridextract.extract(tmp_path / "nothing.xlsx", "Grid", "A1:A1",
                            tmp_path / "work")
    assert "no workbook" in caught.value.message


def test_a_wrong_sheet_name_lists_the_sheets_there_are(tmp_path):
    book = a_workbook(tmp_path / "book.xlsx")
    with pytest.raises(gridextract.CannotRead) as caught:
        gridextract.extract(book, "Assesment", "A1:A1", tmp_path / "work")
    assert "Assesment" in caught.value.message
    assert "Grid" in caught.value.message, "it did not say what sheets exist"


def test_something_that_is_not_a_range_says_what_one_looks_like(tmp_path):
    book = a_workbook(tmp_path / "book.xlsx")
    with pytest.raises(gridextract.CannotRead) as caught:
        gridextract.extract(book, "Grid", "the assessment bit", tmp_path / "work")
    assert "AY98:BK102" in caught.value.message, "no example of a range"


def test_a_range_where_everything_is_hidden_says_so(tmp_path):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Grid"
    sheet.cell(row=1, column=1).value = "hidden"
    sheet.row_dimensions[1].hidden = True
    book.save(str(tmp_path / "book.xlsx"))

    with pytest.raises(gridextract.CannotRead) as caught:
        gridextract.extract(tmp_path / "book.xlsx", "Grid", "A1:A1",
                            tmp_path / "work")
    assert "hidden" in caught.value.message


# ----------------------------------------------------------- the bridge ----
def test_the_bridge_asks_the_same_question_about_windows_that_reveal_does():
    """Two readings of one line is how the update handoff went wrong once."""
    assert office.on_windows() == reveal.on_windows()


def test_no_office_is_a_sentence_and_not_a_crash(tmp_path):
    with pytest.raises(office.OfficeRefused) as caught:
        office.grid_to_image(tmp_path / "b.xlsx", "Grid", "A1:A1",
                             tmp_path / "out.png", using=FakeOffice(here=False))
    assert "not installed" in caught.value.message
    assert "Photographs need no Office" in caught.value.message


def test_the_bridge_imports_nothing_platform_specific_until_it_is_used():
    """Photo pages need no Office at all. A machine without it must keep
    working for everything the app already does."""
    source = Path(office.__file__).read_text()
    head = source.split("def ")[0]
    for forbidden in ("import office_win", "import office_mac", "win32com"):
        assert forbidden not in head, "%s is imported at start up" % forbidden


def test_a_grid_goes_through_the_extraction_before_it_reaches_office(tmp_path):
    """The safety property, asserted rather than trusted: what Office opens is
    the small copy, never the workbook it came from."""
    book = a_workbook(tmp_path / "book.xlsx")
    fake = FakeOffice()
    office.grid_to_image(book, "Grid", "A1:B2", tmp_path / "out.png",
                         workshop=tmp_path / "work", using=fake)
    opened, _ = fake.grids[0]
    assert opened != book, "Office was handed his own workbook"
    assert opened.parent == tmp_path / "work"


def test_the_working_files_are_kept_when_a_workshop_is_named(tmp_path):
    """What the check tool wants when somebody is working out where a chain
    broke."""
    book = a_workbook(tmp_path / "book.xlsx")
    shop = tmp_path / "work"
    office.grid_to_image(book, "Grid", "A1:B2", tmp_path / "out.png",
                         workshop=shop, using=FakeOffice())
    assert (shop / "grid.xlsx").is_file()


def test_the_working_files_are_cleaned_up_when_one_is_not(tmp_path):
    book = a_workbook(tmp_path / "book.xlsx")
    fake = FakeOffice()
    office.grid_to_image(book, "Grid", "A1:B2", tmp_path / "out.png", using=fake)
    opened, _ = fake.grids[0]
    assert not opened.exists(), "a temporary folder was left behind"


def test_office_producing_nothing_is_reported_and_not_passed_on(tmp_path):
    """A zero-length picture would travel silently into a report."""
    book = a_workbook(tmp_path / "book.xlsx")
    with pytest.raises(office.OfficeRefused) as caught:
        office.grid_to_image(book, "Grid", "A1:B2", tmp_path / "out.png",
                             using=FakeOffice(writes=b""))
    assert "produced nothing" in caught.value.message
    assert "Nothing of yours was changed" in caught.value.message


def test_a_missing_document_is_refused_before_word_is_started(tmp_path):
    fake = FakeOffice()
    with pytest.raises(office.OfficeRefused) as caught:
        office.docx_to_pdf(tmp_path / "gone.docx", tmp_path / "out.pdf",
                           using=fake)
    assert "no document" in caught.value.message
    assert not fake.documents, "Word was started for a file that is not there"


def test_a_pdf_that_did_not_appear_is_reported(tmp_path):
    (tmp_path / "one.docx").write_bytes(b"not really a document")
    with pytest.raises(office.OfficeRefused) as caught:
        office.docx_to_pdf(tmp_path / "one.docx", tmp_path / "out.pdf",
                           using=FakeOffice(writes=b""))
    assert "produced nothing" in caught.value.message
    assert "was not changed" in caught.value.message


def test_the_bridge_says_what_is_going_to_do_the_work():
    assert office.describe(using=FakeOffice()) == "a stand-in, not real Office"
    assert "not installed" in office.describe(using=FakeOffice(here=False))


# ------------------------------------------------------- the Mac backend ---
# Real Excel and Word are not touched here. What is testable without them is
# the part most likely to be wrong: the text handed to AppleScript, and what
# happens when Office answers with an error.
import office_mac  # noqa: E402


class Answer:
    """What osascript gave back, so a test can state it."""

    def __init__(self, returncode=0, stderr=""):
        self.returncode = returncode
        self.stderr = stderr
        self.stdout = ""


def test_a_quotation_mark_in_a_path_is_escaped_before_applescript_sees_it():
    """A job folder name is arbitrary text off Mark's disk. Unescaped, a quote
    would end the string early and change what the script says."""
    assert office_mac._quoted('/jobs/the "big" one/x.docx') \
        == '/jobs/the \\"big\\" one/x.docx'


def test_a_backslash_is_escaped_first_so_it_cannot_escape_the_escapes():
    assert office_mac._quoted("a\\b") == "a\\\\b"
    assert office_mac._quoted('a\\"b') == 'a\\\\\\"b'


def test_a_windows_style_path_survives_the_escaping():
    """Not the platform this runs on, and a path can still arrive looking like
    one when it came out of a file somebody wrote elsewhere."""
    assert office_mac._quoted("C:\\jobs\\one.docx") == "C:\\\\jobs\\\\one.docx"


def test_no_office_here_is_a_plain_no(monkeypatch, tmp_path):
    monkeypatch.setattr(office_mac, "EXCEL_APP", tmp_path / "nothing.app")
    assert office_mac.available() is False


def test_word_refusing_says_what_word_said_and_how_to_warm_it(monkeypatch, tmp_path):
    """The fix for a cold Word is to click into it. Quitting it makes things
    worse, so the message says not to."""
    monkeypatch.setattr(office_mac, "_osascript",
                        lambda script, timeout: Answer(1, "execution error: -1708"))
    (tmp_path / "one.docx").write_bytes(b"x")
    with pytest.raises(office.OfficeRefused) as caught:
        office_mac.docx_to_pdf(tmp_path / "one.docx", tmp_path / "out.pdf")
    said = caught.value.message
    assert "-1708" in said, "Word's own words were thrown away"
    assert "Do not quit it" in said


def test_word_claiming_success_with_no_pdf_is_not_believed(monkeypatch, tmp_path):
    monkeypatch.setattr(office_mac, "_osascript",
                        lambda script, timeout: Answer(0, ""))
    (tmp_path / "one.docx").write_bytes(b"x")
    with pytest.raises(office.OfficeRefused) as caught:
        office_mac.docx_to_pdf(tmp_path / "one.docx", tmp_path / "out.pdf")
    assert "it is not there" in caught.value.message


def test_the_path_word_is_given_is_the_escaped_one(monkeypatch, tmp_path):
    awkward = tmp_path / 'a "quoted" job'
    awkward.mkdir()
    document = awkward / "one.docx"
    document.write_bytes(b"x")
    scripts = []

    def watch(script, timeout):
        scripts.append(script)
        return Answer(1, "stopped on purpose")

    monkeypatch.setattr(office_mac, "_osascript", watch)
    with pytest.raises(office.OfficeRefused):
        office_mac.docx_to_pdf(document, tmp_path / "out.pdf")

    written = scripts[0]
    assert 'a \\"quoted\\" job' in written
    assert 'a "quoted" job' not in written, "a raw quote reached AppleScript"


def test_excel_refusing_closes_any_workbook_it_left_open(monkeypatch, tmp_path):
    """A failed attempt can leave the small workbook open, and the next attempt
    would then be answering a dialog instead of doing the work."""
    calls = []

    def watch(script, timeout):
        calls.append(script)
        return Answer(1, "execution error: -50")

    monkeypatch.setattr(office_mac, "_osascript", watch)
    little = tmp_path / "grid.xlsx"
    a_workbook(little)
    with pytest.raises(office.OfficeRefused):
        office_mac.render_grid(little, tmp_path / "out.png")
    assert any("close" in one and "grid.xlsx" in one for one in calls), \
        "nothing tried to close the stray workbook"


# -------------------------------------------------------- the check tool ---
# The tool that proves Office works on a given computer. Its Office steps are
# not exercised here, for the reason at the top of this file. What is testable
# is everything around them, and the folder it writes into.
import office_check  # noqa: E402


def test_the_check_writes_only_into_its_own_folder(monkeypatch, tmp_path):
    """Never near a job, and never anywhere a person keeps work."""
    monkeypatch.setenv("RRF_OFFICE_CHECK", str(tmp_path / "check"))
    assert office_check.where() == tmp_path / "check"


def test_the_practice_workbook_has_the_four_things_that_go_wrong(tmp_path):
    """A merged heading, a filled header, borders, and numbers. A grid that
    renders badly gets one of those wrong, so all four are in the practice
    one or it proves less than it looks."""
    made = office_check.a_practice_workbook(tmp_path)
    sheet = openpyxl.load_workbook(str(made)).active

    assert [str(one) for one in sheet.merged_cells.ranges] == ["A1:D1"]
    assert sheet.cell(row=2, column=1).fill.patternType == "solid"
    assert sheet.cell(row=2, column=1).border.bottom.style
    assert sheet.cell(row=3, column=3).value == 1021630


def test_naming_a_workbook_with_no_sheet_lists_the_sheets(monkeypatch, tmp_path):
    """The likeliest mistake anybody makes running this, so it answers the
    question instead of repeating the usage line."""
    monkeypatch.setenv("RRF_OFFICE_CHECK", str(tmp_path / "check"))
    book = a_workbook(tmp_path / "book.xlsx")
    said = []
    code = office_check.run(book, None, None, out=said.append)
    assert code == 1
    assert "Grid" in "\n".join(said), "it did not say what sheets there are"


def test_the_check_starts_by_saying_what_will_do_the_work(monkeypatch, tmp_path):
    monkeypatch.setenv("RRF_OFFICE_CHECK", str(tmp_path / "check"))
    monkeypatch.setattr(office, "describe", lambda using=None: "a stand-in")
    book = a_workbook(tmp_path / "book.xlsx")
    said = []
    office_check.run(book, None, None, out=said.append)
    assert "a stand-in" in "\n".join(said)


def test_the_check_never_lets_a_traceback_be_the_whole_answer():
    """The person reading it is not going to debug one. A traceback may be
    printed underneath a sentence, never instead of it."""
    source = Path(office_check.__file__).read_text()
    assert "useful" in source
    assert source.count("traceback.format_exc()") == 2


# --------------------------------------------------- the Windows backend ---
# None of this runs on Windows here. What is testable from a Mac is the text
# handed to PowerShell, the numbers Office is given, and the decision about
# which mechanism to use. The Office calls themselves are proven only on the
# virtual machine, and `docs/CHECKS.md` is where that is written down.
import office_win  # noqa: E402


def test_the_windows_backend_imports_on_a_mac():
    """It must, because a test file imports it and because a mistake here
    would only show up on the machine nobody can debug."""
    assert office_win.name()


def test_a_mac_is_not_offered_the_windows_backend():
    assert office_win.available() is False


def test_the_numbers_office_is_given_are_the_documented_ones():
    """Excel's own values for "draw it the way it would print" and "as a
    picture". Wrong numbers give a bitmap or a screen rendering, and both look
    worse than his paste. They cannot be read off Excel before Excel is
    talking, so they are written down and checked here."""
    assert office_win.XL_PRINTER == 2
    assert office_win.XL_PICTURE == -4147
    assert office_win.WD_PDF == 17


def test_a_quote_in_a_path_is_doubled_for_powershell():
    """The only escape a single-quoted PowerShell string has."""
    assert office_win._ps_text("/jobs/o'brien/x.docx") == "/jobs/o''brien/x.docx"


def test_a_path_reaches_powershell_escaped(monkeypatch, tmp_path):
    scripts = []

    def watch(script, work, what):
        scripts.append(script)

    monkeypatch.setattr(office_win, "_run_powershell", watch)
    monkeypatch.setattr(office_win, "_com", lambda: None)
    office_win.docx_to_pdf(tmp_path / "o'brien.docx", tmp_path / "out.pdf")
    assert "o''brien.docx" in scripts[0]


def test_powershell_always_quits_office_even_when_a_step_fails(monkeypatch, tmp_path):
    """An Excel left running holds the file open and the next attempt fails
    for a reason that has nothing to do with the next attempt."""
    scripts = []
    monkeypatch.setattr(office_win, "_run_powershell",
                        lambda script, work, what: scripts.append(script))
    monkeypatch.setattr(office_win, "_com", lambda: None)
    office_win.docx_to_pdf(tmp_path / "one.docx", tmp_path / "out.pdf")
    assert "finally" in scripts[0]
    assert "$word.Quit()" in scripts[0]


def test_excel_is_opened_read_only_and_a_new_copy(monkeypatch, tmp_path):
    """Mark may have a workbook open while the app runs. A new hidden copy
    cannot see his windows and cannot disturb them."""
    scripts = []
    monkeypatch.setattr(office_win, "_run_powershell",
                        lambda script, work, what: scripts.append(script))
    monkeypatch.setattr(office_win, "_com", lambda: None)
    a_workbook(tmp_path / "grid.xlsx")
    office_win.render_grid(tmp_path / "grid.xlsx", tmp_path / "out.png")
    written = scripts[0]
    assert "New-Object -ComObject Excel.Application" in written
    assert "$excel.Visible = $false" in written
    assert "$book = $excel.Workbooks.Open('%s', $false, $true)" \
        % office_win._ps_text(tmp_path / "grid.xlsx") in written, \
        "the workbook was not opened read only"


def test_powershell_is_used_when_the_library_will_not_load(monkeypatch, tmp_path):
    """The whole reason both exist. The package installs libraries in a way
    that never runs their setup steps, so pywin32 may not import at all."""
    monkeypatch.setattr(office_win, "_com", lambda: None)
    used = []
    monkeypatch.setattr(office_win, "_run_powershell",
                        lambda script, work, what: used.append(what))
    office_win.docx_to_pdf(tmp_path / "one.docx", tmp_path / "out.pdf")
    assert used == ["making the PDF"]


def test_the_library_failing_falls_through_to_powershell(monkeypatch, tmp_path):
    """Loading the library and Office answering it are two different things,
    and only the second one matters."""
    class Broken:
        def DispatchEx(self, _what):
            raise OSError("Excel is not answering")

    monkeypatch.setattr(office_win, "_com", lambda: Broken())
    monkeypatch.setattr(office_win, "_note", lambda what, exc: None)
    used = []
    monkeypatch.setattr(office_win, "_run_powershell",
                        lambda script, work, what: used.append(what))
    office_win.docx_to_pdf(tmp_path / "one.docx", tmp_path / "out.pdf")
    assert used == ["making the PDF"], "it stopped instead of falling back"


def test_the_fallback_does_not_hide_that_the_library_failed(monkeypatch, tmp_path):
    """A machine where pywin32 never works would otherwise look normal and
    quietly cost seconds on every grid."""
    class Broken:
        def DispatchEx(self, _what):
            raise OSError("Excel is not answering")

    written = []
    monkeypatch.setattr(office_win, "_com", lambda: Broken())
    monkeypatch.setattr(office_win, "_note",
                        lambda what, exc: written.append((what, str(exc))))
    monkeypatch.setattr(office_win, "_run_powershell",
                        lambda script, work, what: None)
    office_win.docx_to_pdf(tmp_path / "one.docx", tmp_path / "out.pdf")
    assert written and "Excel is not answering" in written[0][1]


def test_waking_the_library_never_raises(monkeypatch):
    """It runs before anything works. A failure here means the fallback is
    used, not that the app stops."""
    def no_libraries():
        raise RuntimeError("nothing is installed")

    monkeypatch.setattr(office_win, "_library_home", no_libraries)
    office_win._wake_pywin32()


def test_no_powershell_at_all_is_a_sentence(monkeypatch, tmp_path):
    monkeypatch.setattr(office_win, "_powershell", lambda: "")
    with pytest.raises(office.OfficeRefused) as caught:
        office_win._run_powershell("x", tmp_path, "making the PDF")
    assert "no PowerShell" in caught.value.message


def test_office_going_quiet_says_to_look_at_office(monkeypatch, tmp_path):
    """Office asks a person questions and cannot tell us it is asking. On
    2026-09-04 Word sat waiting on a Grant access box for seven minutes and
    said nothing. A wait that never ends is the fault we keep fixing."""
    import subprocess as sub

    def times_out(*_a, **_k):
        raise sub.TimeoutExpired("powershell", 600)

    monkeypatch.setattr(office_win, "_powershell", lambda: "/bin/echo")
    monkeypatch.setattr(office_win.subprocess, "run", times_out)
    with pytest.raises(office.OfficeRefused) as caught:
        office_win._run_powershell("x", tmp_path, "making the PDF")
    said = caught.value.message
    assert "did not answer" in said
    assert "waiting for you" in said


def test_the_mac_side_also_says_when_office_goes_quiet(monkeypatch):
    """Same fault, same answer, both platforms. Word going silent on a Grant
    access box is what made this necessary."""
    import subprocess as sub

    def times_out(*_a, **_k):
        raise sub.TimeoutExpired("osascript", 330)

    monkeypatch.setattr(office_mac.subprocess, "run", times_out)
    with pytest.raises(office.OfficeRefused) as caught:
        office_mac._osascript("x", 330)
    assert "did not answer" in caught.value.message
    assert "waiting for you" in caught.value.message


# ------------------------------------------ what the built package carries ---
# The riskiest thing in the Windows half is not the Office calls. It is whether
# pywin32 can import at all, because the package installs libraries with pip's
# --target, which copies files and runs no setup steps. So `pywin32.pth`, the
# file that would wire this up, is copied in and never read.
#
# These run against a real built package, which is the only place the question
# can be asked from a Mac. Whether the wiring then works is Check 13, on the
# virtual machine.
import packaging as apppackaging  # noqa: E402

PACKAGES = Path(__file__).resolve().parents[2] / "build" / "packages"
BUILT = PACKAGES / ("Roy R. Fisher v%s"
                    % (apppackaging.version_of(Path(__file__).resolve().parents[2])
                       or "0.0.0"))
LIBRARIES = BUILT / "program" / "python" / "site-packages"

needs_a_built_package = pytest.mark.skipif(
    not LIBRARIES.is_dir(),
    reason="no package is built; run: python3 tools/package_windows.py")


@needs_a_built_package
@pytest.mark.parametrize("part", [
    "openpyxl",
    "win32com",
    "win32/lib",
    "pywin32_system32/pythoncom314.dll",
    "pywin32_system32/pywintypes314.dll",
])
def test_the_package_carries_what_the_bridge_imports(part):
    """Named one at a time so a missing one says which."""
    assert (LIBRARIES / part).exists(), "%s is not in the package" % part


@needs_a_built_package
def test_the_bootstrap_finds_the_folders_in_a_real_package(monkeypatch):
    """The path work, proven against a real package rather than a guess.

    What this cannot prove is that Windows then loads the two system files.
    That is Check 13 and it needs the virtual machine.
    """
    monkeypatch.setattr(office_win, "_library_home", lambda: LIBRARIES)
    monkeypatch.setattr(sys, "path", list(sys.path))
    before = list(sys.path)
    office_win._wake_pywin32()
    added = [one for one in sys.path if one not in before]
    assert str(LIBRARIES / "win32") in added
    assert str(LIBRARIES / "win32" / "lib") in added, "pythoncom lives here"


@needs_a_built_package
def test_the_check_tool_ships():
    """It is the only way anybody finds out whether Office works on a given
    computer, and it is useless if it is not in the package."""
    engine = BUILT / "program" / "app" / "engine"
    for one in ("office.py", "office_win.py", "office_mac.py",
                "gridextract.py", "office_check.py"):
        assert (engine / one).is_file(), "%s is not in the package" % one
