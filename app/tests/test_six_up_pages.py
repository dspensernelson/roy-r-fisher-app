"""Six photographs to a page.

The page arithmetic here is proved with synthetic images and synthetic
captions, which is valid evidence for exactly what it claims: how many pages
a count produces, which rows survive trimming, and where a caption lands.
It makes no claim about Mark's layouts. The one test that does claim
something about his work reads his delivered captions and carries the
corpus marker.
"""
import json
import math
import sys
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Emu
from PIL import Image

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "engine"))
from photo_pages import SIX_UP, build_photo_docx  # noqa: E402

from conftest import CORPUS, has_template  # noqa: E402

SIX_UP_TEMPLATE = Path(__file__).resolve().parents[1] / "templates" / "Photo six-up.docx"

# A landscape photograph, which is what every one of Mark's is.
LANDSCAPE = (400, 300)


def a_job(tmp_path: Path, count: int, caption: str = None) -> Path:
    names = []
    for i in range(count):
        p = tmp_path / f"IMG_{5100 + i}.jpg"
        Image.new("RGB", LANDSCAPE, (i * 30 % 255, 80, 90)).save(p)
        names.append(p.name)
    m = tmp_path / "photo-manifest.json"
    m.write_text(json.dumps({"photos": [
        {"file": n, "caption": caption if caption is not None else f"View of test subject {i}"}
        for i, n in enumerate(names)]}))
    return m


def images_in(cell) -> int:
    return len(cell._tc.findall(".//" + qn("a:blip")))


# --- how many pages -------------------------------------------------------

@pytest.mark.parametrize("count,pages", [(1, 1), (5, 1), (6, 1), (7, 2), (12, 2), (13, 3)])
def test_six_photographs_fill_one_page(tmp_path, count, pages):
    out = build_photo_docx(a_job(tmp_path, count), SIX_UP_TEMPLATE, layout=SIX_UP)
    assert len(Document(str(out)).tables) == pages


@pytest.mark.parametrize("count", [1, 2, 3, 4, 5, 6, 7, 11, 12, 13])
def test_every_photograph_reaches_the_document(tmp_path, count):
    out = build_photo_docx(a_job(tmp_path, count), SIX_UP_TEMPLATE, layout=SIX_UP)
    d = Document(str(out))
    assert len(d.inline_shapes) == count
    captions = [c.text.strip() for t in d.tables for r in t.rows
                for c in r.cells if c.text.strip()]
    assert len(captions) == count
    assert captions == [f"View of test subject {i}" for i in range(count)]


@pytest.mark.parametrize("count", [1, 2, 3, 6, 7, 12, 13])
def test_no_trailing_blank_page(tmp_path, count):
    """The defect that produced an extra page on three-up was the empty
    paragraph after a full last table. Six-up declares its rows honestly so
    the rows do not grow past the page, but the trailing paragraph is
    removed by the same code and this is what proves it still runs."""
    out = build_photo_docx(a_job(tmp_path, count), SIX_UP_TEMPLATE, layout=SIX_UP)
    d = Document(str(out))
    body = list(d.element.body)
    last_table = max(i for i, el in enumerate(body)
                     if el.tag.split("}")[1] == "tbl")
    after = [el for el in body[last_table + 1:]
             if el.tag.split("}")[1] == "p" and "".join(el.itertext()).strip() == ""]
    assert after == [], "an empty paragraph after the last table is an extra page"


# --- the last page --------------------------------------------------------

@pytest.mark.parametrize("count,rows", [(1, 2), (2, 2), (3, 4), (4, 4), (5, 6), (6, 6)])
def test_the_last_page_keeps_whole_pairs(tmp_path, count, rows):
    out = build_photo_docx(a_job(tmp_path, count), SIX_UP_TEMPLATE, layout=SIX_UP)
    assert len(Document(str(out)).tables[-1].rows) == rows


@pytest.mark.parametrize("count", [1, 3, 5])
def test_an_odd_photograph_leaves_exactly_one_empty_box(tmp_path, count):
    """Spenser chose this on 2026-09-07 over merging the last pair into one
    full-width cell that spans the page.

    His reason: a photograph that changes size according to how many
    photographs there happen to be is a rule that surprises somebody six
    months later. So the odd one sits in the left column at the same size as
    every other photograph, and the box beside it is empty and stays bordered.
    This is a decision, not a gap, and that is why it has a test.
    """
    out = build_photo_docx(a_job(tmp_path, count), SIX_UP_TEMPLATE, layout=SIX_UP)
    last = Document(str(out)).tables[-1]
    blank = [c for r in last.rows for c in r.cells
             if not c.text.strip() and not images_in(c)]
    assert len(blank) == 2      # the photograph's cell, and its caption's cell


def test_an_even_photograph_leaves_no_empty_box(tmp_path):
    out = build_photo_docx(a_job(tmp_path, 4), SIX_UP_TEMPLATE, layout=SIX_UP)
    last = Document(str(out)).tables[-1]
    blank = [c for r in last.rows for c in r.cells
             if not c.text.strip() and not images_in(c)]
    assert blank == []


# --- where a photograph actually lands ------------------------------------

def test_reading_order_is_left_to_right_then_down(tmp_path):
    """Spenser, 2026-09-07. The captions must come out in the order he typed
    them, read the way a person reads a page."""
    out = build_photo_docx(a_job(tmp_path, 6), SIX_UP_TEMPLATE, layout=SIX_UP)
    t = Document(str(out)).tables[0]
    assert [t.rows[r].cells[c].text.strip() for r, c in
            ((1, 0), (1, 1), (3, 0), (3, 1), (5, 0), (5, 1))] == \
           [f"View of test subject {i}" for i in range(6)]
    for r, c in ((0, 0), (0, 1), (2, 0), (2, 1), (4, 0), (4, 1)):
        assert images_in(t.rows[r].cells[c]) == 1, (r, c)


def test_a_photograph_is_fitted_to_the_six_up_box(tmp_path):
    """3.20 x 2.40in, which is 4:3 at the six-up column width. Two of them
    share the same 6.70in the one three-up photograph had."""
    out = build_photo_docx(a_job(tmp_path, 1), SIX_UP_TEMPLATE, layout=SIX_UP)
    shape = Document(str(out)).inline_shapes[0]
    assert round(Emu(shape.width).inches, 2) == 3.20
    assert round(Emu(shape.height).inches, 2) == 2.40


# --- the vertical budget --------------------------------------------------

LINE_IN = 0.20          # a 12pt Times line
CHARS_PER_LINE = 48     # measured: a 95-character caption takes 2 lines at 3.20in


def caption_height(text: str) -> float:
    """One caption's height: the declared 0.46in, or taller if it wraps far.

    Deliberately arithmetic rather than a font measurement. Reading a real
    font file would tie this test to one machine's fonts, and the claim here
    is about gross overflow, not typesetting.
    """
    lines = max(1, math.ceil(len(text) / CHARS_PER_LINE))
    return max(0.46, lines * LINE_IN + 0.06)


def page_inches_used(captions) -> float:
    """What one full six-up page occupies, heading included.

    **A page has three caption rows, not six.** Each row carries two captions
    side by side and is as tall as the taller of them. Getting that wrong is
    what made the first cut of this test claim Mark's own worst page
    overflowed by an inch when the built document fits with 0.40in to spare.
    """
    pairs = [captions[i:i + 2] for i in range(0, len(captions), 2)]
    return 0.42 + 3 * 2.60 + sum(max(caption_height(c) for c in pair)
                                 for pair in pairs)


USABLE_IN = 10.0        # 11.0 page less 0.5 top and 0.5 bottom margin


def test_a_full_page_of_ordinary_captions_fits_with_slack_to_spare():
    ordinary = ["View of the northwest corner facing southeast"] * 6
    used = page_inches_used(ordinary)
    assert used <= USABLE_IN
    assert round(USABLE_IN - used, 2) == 0.40


def test_the_page_holds_two_over_long_caption_rows_and_no_more():
    """The limit, written down so nobody meets it by surprise.

    0.40in of slack is two extra caption lines, and a line is spent per ROW,
    not per caption: two over-long captions sitting beside each other in the
    same row cost one line between them. A third over-long row overflows.

    This is a characterisation test. It records where the edge is, it does
    not approve of it. What the app should do when a job crosses it is
    undecided and is the open question in
    docs/plans/2026-09-07-six-up-photo-pages.md.
    """
    short = "View of the west entrance"
    long = "x" * 120                       # three lines at 3.20in

    # Side by side in one row: one row grows, the page is comfortable.
    assert page_inches_used([long, long] + [short] * 4) <= USABLE_IN
    # Two rows grow: exactly at the limit.
    assert page_inches_used([long, short, long, short, short, short]) <= USABLE_IN
    # Three rows grow: past it.
    assert page_inches_used([long, short, long, short, long, short]) > USABLE_IN


@has_template
def test_the_worst_page_marks_real_captions_can_produce_still_fits():
    """The six longest captions in every report he has delivered.

    This is the claim that matters, and it is the one synthetic text cannot
    support: the edge above exists, and his actual writing does not reach it.
    """
    seen = set()
    for doc in sorted(CORPUS.rglob("*.docx")):
        if doc.name.startswith("~$") or "photo" not in doc.name.lower():
            continue
        try:
            d = Document(str(doc))
        except Exception:
            continue
        for t in d.tables:
            for r in t.rows:
                if len(r.cells) < 2:
                    continue
                text = r.cells[1].text.strip()
                if len(text) > 3:
                    seen.add(text)
    assert len(seen) > 100, "the corpus should hold hundreds of real captions"
    worst = sorted(seen, key=len, reverse=True)[:6]
    assert page_inches_used(worst) <= USABLE_IN
